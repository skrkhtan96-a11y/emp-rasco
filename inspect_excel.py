import sys
import os

try:
    import openpyxl
except ImportError:
    print("Installing openpyxl...")
    os.system("pip install openpyxl -q")
    import openpyxl

# Try all xlsx files in Downloads by size (biggest first = most likely employees list)
downloads = r"C:\Users\TWc\Downloads"
xlsx_files = []
for f in os.listdir(downloads):
    if f.endswith('.xlsx'):
        full = os.path.join(downloads, f)
        xlsx_files.append((os.path.getsize(full), full, f))

xlsx_files.sort(reverse=True)

print("=== ALL XLSX FILES (largest first) ===")
for size, path, name in xlsx_files[:10]:
    print(f"  {size:,} bytes  {name}")

print("\n=== INSPECTING TOP 3 XLSX FILES ===\n")

for size, path, name in xlsx_files[:3]:
    print(f"\n{'='*60}")
    print(f"FILE: {name}  ({size:,} bytes)")
    print(f"{'='*60}")
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        print(f"Sheets: {wb.sheetnames}")
        
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            print(f"\n  Sheet: '{sheet_name}'  (max_row={ws.max_row}, max_col={ws.max_column})")
            
            if ws.max_row and ws.max_row > 0:
                # Read header row
                headers = []
                for row in ws.iter_rows(min_row=1, max_row=1, values_only=True):
                    headers = [str(c) if c is not None else '' for c in row]
                
                print(f"  Headers ({len(headers)} columns): {headers}")
                
                # Read first 3 data rows
                row_count = 0
                for row in ws.iter_rows(min_row=2, max_row=4, values_only=True):
                    row_data = [str(c) if c is not None else '' for c in row]
                    print(f"  Row {row_count+2}: {row_data[:15]}")  # first 15 cols
                    row_count += 1
                
                # If more than 100 rows, this is likely the employees sheet
                if ws.max_row and ws.max_row > 100:
                    print(f"\n  *** THIS LOOKS LIKE THE EMPLOYEES SHEET ({ws.max_row} rows) ***")
                    
                    # Find Location/Region column
                    loc_col_idx = None
                    proj_col_idx = None
                    emp_name_idx = None
                    emp_id_idx = None
                    emp_no_idx = None
                    
                    for i, h in enumerate(headers):
                        h_lower = h.lower().strip()
                        if any(x in h_lower for x in ['location', 'region', 'منطقة', 'موقع']):
                            loc_col_idx = i
                            print(f"  -> Location/Region column: [{i}] '{h}'")
                        if any(x in h_lower for x in ['project', 'مشروع']):
                            proj_col_idx = i
                            print(f"  -> Project column: [{i}] '{h}'")
                        if any(x in h_lower for x in ['name', 'employee name', 'اسم', 'full name']):
                            if emp_name_idx is None:
                                emp_name_idx = i
                                print(f"  -> Employee Name column: [{i}] '{h}'")
                        if any(x in h_lower for x in ['emp no', 'emp#', 'employee number', 'employee no', 'empno', 'رقم الموظف']):
                            emp_no_idx = i
                            print(f"  -> Employee Number column: [{i}] '{h}'")
                        if any(x in h_lower for x in ['id', 'employee id', 'emp id', 'رقم']):
                            if emp_id_idx is None and 'name' not in h_lower:
                                emp_id_idx = i
                                print(f"  -> Employee ID column: [{i}] '{h}'")
                    
                    # Extract distinct locations
                    if loc_col_idx is not None:
                        locations = set()
                        projects_by_loc = {}
                        all_projects = set()
                        
                        row_idx = 0
                        for row in ws.iter_rows(min_row=2, values_only=True):
                            val = row[loc_col_idx] if loc_col_idx < len(row) else None
                            prj = row[proj_col_idx] if proj_col_idx is not None and proj_col_idx < len(row) else None
                            
                            if val:
                                loc_str = str(val).strip()
                                locations.add(loc_str)
                                if prj:
                                    prj_str = str(prj).strip()
                                    if loc_str not in projects_by_loc:
                                        projects_by_loc[loc_str] = set()
                                    projects_by_loc[loc_str].add(prj_str)
                                    all_projects.add(prj_str)
                            row_idx += 1
                        
                        print(f"\n  TOTAL DATA ROWS: {row_idx}")
                        print(f"\n  DISTINCT LOCATIONS ({len(locations)}):")
                        for loc in sorted(locations):
                            count = sum(1 for row in ws.iter_rows(min_row=2, values_only=True) 
                                      if row[loc_col_idx] and str(row[loc_col_idx]).strip() == loc)
                            print(f"    '{loc}': {len(projects_by_loc.get(loc, set()))} projects")
                        
                        print(f"\n  REGION → PROJECTS MAPPING:")
                        for loc in sorted(projects_by_loc.keys()):
                            print(f"    '{loc}' → {sorted(projects_by_loc[loc])}")
                        
                        print(f"\n  ALL DISTINCT PROJECTS ({len(all_projects)}):")
                        for p in sorted(all_projects):
                            print(f"    '{p}'")
        
        wb.close()
    except Exception as e:
        print(f"  ERROR: {e}")

print("\n=== DONE ===")
